"""Build the workshop slide deck from editable PowerPoint text and images.

Run from the repository root with the pyenv-created .venv:
    .venv/bin/python slides/build_deck.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "slides" / "cvs-redis-iris-workshop.pptx"
ASSETS = ROOT / "slides" / "assets"

RED = RGBColor(255, 68, 56)
INK = RGBColor(19, 30, 38)
MUTED = RGBColor(83, 103, 113)
WHITE = RGBColor(255, 255, 255)
PALE = RGBColor(245, 248, 249)


def add_text(slide, text, x, y, w, h, size=22, color=INK, bold=False, font="Arial"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0)
    frame.margin_top = frame.margin_bottom = Inches(0)
    for i, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.space_after = Pt(9)
        for run in paragraph.runs:
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
    return box


def new_slide(deck, title, number, notes="", dark=False):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = INK if dark else WHITE
    foreground = WHITE if dark else INK
    add_text(slide, title, 0.65, 0.42, 11.9, 0.72, 33, foreground, True)
    add_text(slide, f"CVS onsite Redis Iris workshop  /  {number:02d}", 0.65, 7.05, 11.9, 0.25, 11, MUTED if not dark else PALE)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def label(slide, text, x=0.65, y=1.35, color=RED):
    add_text(slide, text.upper(), x, y, 11.9, 0.28, 12, color, True)


def body(slide, text, x=0.65, y=1.8, w=11.9, h=4.8, size=23, color=INK):
    add_text(slide, text, x, y, w, h, size, color)


def code(slide, text, x, y, w, h, size=18, color=INK):
    add_text(slide, text, x, y, w, h, size, color, font="Menlo")


deck = Presentation()
deck.slide_width = Inches(13.333)
deck.slide_height = Inches(7.5)

# 1. Cover
s = new_slide(deck, "Redis Iris workshop for CVS", 1, dark=True)
add_text(s, "Four labs. Python only. Fictional retail data.", 0.68, 1.74, 9.9, 0.7, 27, WHITE)
add_text(s, "Vector search\nContext Retriever\nLangCache\nAgent Memory", 0.68, 3.1, 8.5, 2.7, 24, WHITE)
s.shapes.add_picture(str(ASSETS / "redis-logo.png"), Inches(10.0), Inches(2.4), width=Inches(2.2))
s.notes_slide.notes_text_frame.text = (
    "Redis logo source: Redis Iris Retail workshop repository, "
    "https://github.com/Redislabs-Solution-Architects/redis-iris-workshop"
)

# 2. Order and output
s = new_slide(deck, "Workshop path", 2)
label(s, "Work from the repository root")
body(s, "1   Vector search: policy documents                 30 min\n"
        "2   Context Retriever: live orders                 45 min\n"
        "3   LangCache: a public policy answer              20 min\n"
        "4   Agent Memory: pickup preference                25 min", y=1.9, size=23)
add_text(s, "Allow 20 min for Cloud and Python setup, then 10 min to explain what you built. Total: 150 min.",
         0.65, 6.1, 11.5, 0.72, 18, MUTED)

# 3. Free Cloud account
s = new_slide(deck, "Redis Cloud account and free database", 3,
              "Source: https://redis.io/docs/latest/operate/rc/databases/create-database/create-free-database/")
label(s, "Cloud setup")
body(s, "1   Sign up or sign in at cloud.redis.io\n"
        "2   Select New database\n"
        "3   Choose Try 30 MB for free\n"
        "4   Choose vendor and region, then Create database", y=1.85, size=25)
add_text(s, "Redis allows one free database per account.", 0.65, 6.15, 11.2, 0.4, 19, MUTED)

# 4. Database credentials and screenshot from reference workshop
s = new_slide(deck, "Database endpoint and password", 4,
              "Sources: https://redis.io/docs/latest/operate/rc/databases/connect/ ; "
              "Illustrative screenshot from https://github.com/Redislabs-Solution-Architects/redis-iris-workshop/blob/main/credentials.png . "
              "Console layout can change.")
label(s, "Database configuration")
body(s, "Copy the public endpoint: host and port.\n\n"
        "Reveal the Default user password.\n\n"
        "The Connect wizard shows a Python redis-py example.\n\n"
        "Put the values in REDIS_URL inside .env.", x=0.65, y=1.85, w=6.15, h=4.75, size=22)
s.shapes.add_picture(str(ASSETS / "redis-cloud-credentials.png"), Inches(7.45), Inches(1.28),
                     width=Inches(4.4), height=Inches(5.83))

# 5. Local environment
s = new_slide(deck, "Local Python setup", 5)
label(s, "Pyenv-selected Python")
code(s, "pyenv exec python3 -m venv .venv\n"
        "source .venv/bin/activate\n"
        "python -m pip install -r requirements.txt\n"
        "cp .env.example .env", 0.67, 1.85, 11.7, 2.35, 22)
add_text(s, "Fill REDIS_URL, pre-download the local Hugging Face model, then run the README ping check.",
         0.67, 5.04, 11.6, 1.0, 22, INK)
add_text(s, "Use rediss:// when the database requires TLS. Keep .env out of Git.",
         0.67, 6.14, 11.6, 0.48, 18, MUTED)

# 6. Services and eligibility
s = new_slide(deck, "Cloud services used by the labs", 6,
              "Sources: https://redis.io/docs/latest/operate/rc/langcache/create-service/ ; "
              "https://redis.io/docs/latest/develop/ai/context-engine/context-retriever/quickstart/ ; "
              "https://redis.io/docs/latest/operate/iris/agent-memory/create-service/")
label(s, "Database plus managed services")
body(s, "Lab 1: RedisVL vector search on the database\n\n"
        "Lab 2: Context Retriever surface over live HASHes\n\n"
        "Lab 3: LangCache service, cache ID, service key\n\n"
        "Lab 4: Agent Memory service, store ID, service key", y=1.85, size=21)
add_text(s, "Agent Memory Quick create can use Free 30 MB. Confirm service access and keep an instructor service ready.",
         0.65, 5.95, 11.8, 0.95, 19, RED, True)

# 7. Lab 1
s = new_slide(deck, "Lab 1: vector search", 7,
              "Sources: https://docs.redisvl.com/en/latest/user_guide/04_vectorizers.html ; "
              "https://docs.redisvl.com/en/latest/api/searchindex.html . "
              "Workshop guide: exercises/lab_1/README.md")
label(s, "Edit exercises/lab_1/01_vector.py")
body(s, "Read eight policy Markdown files.\n\n"
        "Embed its body locally with all-MiniLM-L6-v2.\n\n"
        "Create a RedisVL HASH index and load records.\n\n"
        "Inspect the 1,536-byte vector; query top 2.", x=0.65, y=1.8, w=7.25, h=4.95, size=22)
code(s, "python exercises/lab_1/01_vector.py\n\nExpected: pickup policy near top", 8.1, 2.12, 4.45, 2.6, 17)

# 8. Lab 2 data/model
s = new_slide(deck, "Lab 2: data and entity model", 8,
              "Source: https://redis.io/docs/latest/develop/ai/context-engine/context-retriever/quickstart/ . "
              "Workshop guide: exercises/lab_2/README.md")
label(s, "Edit the loader and models.py")
body(s, "Load stores.jsonl and orders.jsonl as HASHes.\n\n"
        "Inspect O1001 with HGETALL.\n\n"
        "Keep keys aligned with __redis_key_template__.\n\n"
        "Mark entity IDs as key components.\n\n"
        "Index status as a tag.", x=0.65, y=1.8, w=7.7, h=4.9, size=21)
code(s, "workshop:store:S101\nworkshop:order:O1001\n\nExpected: Stores 2, Orders 3", 8.5, 2.18, 4.05, 2.6, 17)

# 9. Context Retriever setup
s = new_slide(deck, "Context Retriever surface and keys", 9,
              "Source: https://redis.io/docs/latest/develop/ai/context-engine/context-retriever/quickstart/")
label(s, "Cloud sign-in and ctxctl")
body(s, "Sign in with ctxctl. Create an admin key.\n\n"
        "Create a surface from exercises/lab_2/models.py and your Redis endpoint.\n\n"
        "Create an agent key for tool calls.", x=0.65, y=1.85, w=7.0, h=4.6, size=22)
code(s, "ctxctl auth login -u EMAIL\n"
        "ctxctl admin create --name cvs-workshop-admin\n"
        "ctxctl surface create --models exercises/lab_2/models.py ...\n"
        "ctxctl agent create --surface-id ID ...", 7.9, 2.0, 4.7, 3.55, 15)
add_text(s, "Admin key manages the surface. Agent key invokes generated tools.", 0.65, 6.08, 11.7, 0.65, 18, MUTED)

# 10. Lab 2 tool call
s = new_slide(deck, "Lab 2: generated order tool", 10,
              "Source: https://redis.io/tutorials/getting-started-with-redis-iris/ . "
              "Workshop guide: exercises/lab_2/README.md")
label(s, "Edit exercises/lab_2/02_context.py")
body(s, "Print every generated tool with the agent key.\n\n"
        "Inspect get_order_by_id and its input schema.\n\n"
        "Call it with id = O1001.\n\n"
        "Change the sample order, reload, and call again.", x=0.65, y=1.8, w=7.8, h=4.9, size=22)
code(s, "await client.list_tools(agent_key)\n\n"
        "await client.query_tool(\n"
        "  agent_key=agent_key,\n"
        "  tool_name='get_order_by_id',\n"
        "  arguments={'id': 'O1001'})", 8.35, 1.95, 4.35, 3.45, 16)

# 11. Lab 3
s = new_slide(deck, "Lab 3: LangCache", 11,
              "Sources: https://redis.io/docs/latest/operate/rc/langcache/create-service/ ; "
              "https://redis.io/docs/latest/operate/rc/langcache/use-langcache/ ; "
              "https://docs.redisvl.com/en/latest/user_guide/13_langcache_semantic_cache.html . "
              "Workshop guide: exercises/lab_3/README.md")
label(s, "Create service in Redis Cloud")
body(s, "Open LangCache and create a service on the workshop database.\n\n"
        "Copy the API key once. Get the base URL and cache ID from Connectivity.\n\n"
        "Store one answer, then check the same rephrasing at three distance thresholds.",
        x=0.65, y=1.8, w=8.0, h=4.85, size=21)
code(s, "for limit in (0.02, 0.10, 0.40):\n"
        "  hits = cache.check(\n"
        "    prompt=similar,\n"
        "    distance_threshold=limit)\n"
        "  print(limit, bool(hits))", 8.3, 2.16, 4.4, 3.0, 15)

# 12. Lab 4
s = new_slide(deck, "Lab 4: Agent Memory", 12,
              "Source: https://redis.io/docs/latest/develop/ai/context-engine/agent-memory/python-sdk-quickstart/ . "
              "Workshop guide: exercises/lab_4/README.md")
label(s, "Create an Agent Memory service")
body(s, "Use Quick create on Free, or set a 1-minute cadence in a custom service.\n\n"
        "Copy its API key, endpoint, and store ID. Set a unique WORKSHOP_USER_ID.\n\n"
        "Check service health, write one fictional event, then compare session and long-term memory.",
        x=0.65, y=1.8, w=8.0, h=4.7, size=21)
code(s, "python exercises/lab_4/04_memory.py write\n\n"
        "python exercises/lab_4/04_memory.py read", 7.85, 2.2, 4.85, 2.9, 15)
add_text(s, "The session appears immediately. Long-term extraction runs later.",
         0.65, 6.12, 11.8, 0.65, 18, MUTED)

# 13. Debrief
s = new_slide(deck, "Debrief: choose the right source", 13)
label(s, "Explain the system behind each answer")
body(s, "What does the pickup policy say?\n\n"
        "What is order O1001's status right now?\n\n"
        "Have we answered this public question before?\n\n"
        "Which store does this fictional user prefer?", y=1.75, size=24)
add_text(s, "For each answer: what was stored, who updates it, and can it become stale?",
         0.65, 6.14, 11.8, 0.62, 18, MUTED)

OUT.parent.mkdir(exist_ok=True)
deck.save(OUT)
print(OUT)
